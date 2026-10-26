"""Escritura de los CSV y reportes por consola."""

import csv
from datetime import datetime
from pathlib import Path

from datos import formatear_minutos, formatear_pesos

CAMPOS_LIQUIDACION = [
    "id", "patente", "tipo_vehiculo", "entrada", "salida", "dias", "minutos",
    "bloques_30", "horas_adicionales", "estacionamiento_base", "bonificacion",
    "estacionamiento", "solicito_lavado", "lavado", "total", "tope_aplicado",
]

CAMPOS_RECHAZADOS = ["id", "patente", "tipo_vehiculo", "motivo"]

CAMPOS_MANUAL = ["id", "fecha_registro"] + [
    campo for campo in CAMPOS_LIQUIDACION if campo != "id"
]


def ordenar_liquidacion(validos):
    """Ordena la liquidacion por patente y, si hay empatados, por ID."""
    return sorted(validos, key=lambda fila: (fila["patente"], fila["id"]))


def ordenar_rechazados(rechazados):
    """Ordena los rechazos por ID para que sean faciles de ubicar."""
    return sorted(rechazados, key=lambda fila: (fila["id"], fila["patente"]))


def contar_rechazos(rechazados):
    """Devuelve un diccionario motivo -> cantidad, ordenado por motivo."""
    conteo = {}
    for fila in rechazados:
        conteo[fila["motivo"]] = conteo.get(fila["motivo"], 0) + 1
    return dict(sorted(conteo.items()))


def escribir_csv(ruta, campos, filas):
    """Escribe un CSV completo (borra el anterior)."""
    with Path(ruta).open("w", newline="", encoding="utf-8-sig") as archivo:
        escritor = csv.DictWriter(archivo, fieldnames=campos)
        escritor.writeheader()
        for fila in filas:
            escritor.writerow({campo: fila.get(campo, "") for campo in campos})


def escribir_liquidacion(validos, ruta):
    escribir_csv(ruta, CAMPOS_LIQUIDACION, ordenar_liquidacion(validos))


def escribir_rechazados(rechazados, ruta):
    escribir_csv(ruta, CAMPOS_RECHAZADOS, ordenar_rechazados(rechazados))


def escribir_ranking(validos, ruta):
    """Desafio nivel 1 N.º 3: mismos registros, de mayor a menor total."""
    ordenados = sorted(validos, key=lambda fila: (-fila["total"], fila["patente"]))
    escribir_csv(ruta, CAMPOS_LIQUIDACION, ordenados)


def agregar_liquidacion_manual(registro, ruta):
    """Agrega una fila a liquidacion-manual.csv sin borrar lo anterior."""
    ruta = Path(ruta)
    existe = ruta.exists() and ruta.stat().st_size > 0

    if existe:
        archivo = ruta.open("a", newline="", encoding="utf-8")
    else:
        # utf-8-sig solo al crear: si se usa al agregar, mete un BOM en el medio.
        archivo = ruta.open("w", newline="", encoding="utf-8-sig")

    with archivo:
        escritor = csv.DictWriter(archivo, fieldnames=CAMPOS_MANUAL)
        if not existe:
            escritor.writeheader()
        escritor.writerow({campo: registro.get(campo, "") for campo in CAMPOS_MANUAL})


def imprimir_desglose(patente, tipo_vehiculo, liquidacion):
    print()
    print("=" * 52)
    print("LIQUIDACION")
    print("=" * 52)
    print("Patente:  {}".format(patente))
    print("Tipo:     {}".format(tipo_vehiculo))

    if liquidacion["dias"] > 0:
        print("Tiempo:   {} dia(s) completo(s)".format(liquidacion["dias"]))
    else:
        detalle = "{} ({} bloques de 30 min".format(
            formatear_minutos(liquidacion["minutos"]), liquidacion["bloques_30"])
        if liquidacion["horas_adicionales"]:
            detalle += " + {} hora(s) adicional(es)".format(
                liquidacion["horas_adicionales"])
        detalle += ")"
        print("Tiempo:   " + detalle)

    print("Estacionamiento base:      {:>12}".format(
        formatear_pesos(liquidacion["estacionamiento_base"])))
    if liquidacion["bonificacion"]:
        print("Bonificacion por lavado:  {:>12}".format(
            "-" + formatear_pesos(liquidacion["bonificacion"])))
    print("Estacionamiento a cobrar:  {:>12}".format(
        formatear_pesos(liquidacion["estacionamiento"])))
    print("Lavado:                    {:>12}".format(
        formatear_pesos(liquidacion["lavado"])))
    if liquidacion["tope_aplicado"]:
        print("(se aplico el tope diario de la categoria)")
    print("-" * 52)
    print("TOTAL A PAGAR:             {:>12}".format(
        formatear_pesos(liquidacion["total"])))
    print("=" * 52)


def imprimir_resumen(validos, rechazados):
    total = sum(fila["total"] for fila in validos)
    print()
    print("=" * 52)
    print("RESUMEN DEL LOTE")
    print("=" * 52)
    print("Procesados:      {}".format(len(validos)))
    print("Rechazados:      {}".format(len(rechazados)))
    print("Total recaudado: {}".format(formatear_pesos(total)))

    if rechazados:
        print()
        print("Rechazos por motivo:")
        for motivo, cantidad in contar_rechazos(rechazados).items():
            print("  - {}: {}".format(motivo, cantidad))


def calcular_metricas(validos):
    """Desafio nivel 1 N.º 2: datos para el reporte ejecutivo."""
    recaudado = {}
    cantidad = {}
    minutos = {}
    lavados = 0

    for fila in validos:
        tipo = fila["tipo_vehiculo"]
        recaudado[tipo] = recaudado.get(tipo, 0) + fila["total"]
        cantidad[tipo] = cantidad.get(tipo, 0) + 1

        # Las estadias prolongadas se miden en dias: se pasan a minutos para
        # que el promedio sea comparable con las del mismo dia.
        duracion = fila["dias"] * 1440 if fila["dias"] > 0 else fila["minutos"]
        minutos[tipo] = minutos.get(tipo, 0) + duracion

        if fila["solicito_lavado"] == "S":
            lavados += 1

    promedios = {tipo: minutos[tipo] / cantidad[tipo] for tipo in cantidad}
    mas_rentable = max(recaudado, key=recaudado.get) if recaudado else None
    porcentaje_lavado = (lavados / len(validos) * 100) if validos else 0

    return {
        "recaudado": recaudado,
        "cantidad": cantidad,
        "promedios": promedios,
        "mas_rentable": mas_rentable,
        "porcentaje_lavado": porcentaje_lavado,
        "lavados": lavados,
        "total_vehiculos": len(validos),
    }


def imprimir_metricas(metricas):
    print()
    print("-" * 52)
    print("METRICAS DEL NEGOCIO (desafio nivel 1 N.º 2)")
    print("-" * 52)

    if metricas["mas_rentable"] is None:
        print("No hay movimientos validos para calcular metricas.")
        return

    tipo = metricas["mas_rentable"]
    print("Categoria mas rentable: {} ({})".format(
        tipo, formatear_pesos(metricas["recaudado"][tipo])))
    print("Vehiculos con lavado:   {} de {} ({:.1f}%)".format(
        metricas["lavados"], metricas["total_vehiculos"],
        metricas["porcentaje_lavado"]).replace(".", ","))
    print("Estadia promedio por categoria:")
    for categoria in sorted(metricas["promedios"]):
        print("  - {}: {:.1f} min".format(
            categoria, metricas["promedios"][categoria]).replace(".", ","))

    print()
    print("Recaudacion por categoria:")
    for categoria in sorted(metricas["recaudado"]):
        print("  - {}: {} ({} vehiculos)".format(
            categoria,
            formatear_pesos(metricas["recaudado"][categoria]),
            metricas["cantidad"][categoria]))


# ---------------------------------------------------------------------------
# Reporte Excel (opcional: usa openpyxl; el resto del sistema no lo necesita)
# ---------------------------------------------------------------------------

CAMPOS_RANKING = ["puesto"] + CAMPOS_LIQUIDACION

ANCHOS_LIQUIDACION = {
    "id": 22,
    "patente": 12,
    "tipo_vehiculo": 14,
    "entrada": 9,
    "salida": 9,
    "dias": 7,
    "minutos": 9,
    "bloques_30": 10,
    "horas_adicionales": 12,
    "estacionamiento_base": 17,
    "bonificacion": 14,
    "estacionamiento": 17,
    "solicito_lavado": 12,
    "lavado": 12,
    "total": 14,
    "tope_aplicado": 12,
    "puesto": 8,
}

ANCHOS_RECHAZADOS = {
    "id": 22,
    "patente": 14,
    "tipo_vehiculo": 16,
    "motivo": 48,
}

FORMATOS_LIQUIDACION = {
    "dias": "0",
    "minutos": "0",
    "bloques_30": "0",
    "horas_adicionales": "0",
    "estacionamiento_base": '"$"#,##0',
    "bonificacion": '"$"#,##0',
    "estacionamiento": '"$"#,##0',
    "lavado": '"$"#,##0',
    "total": '"$"#,##0',
    "puesto": "0",
}


def escribir_excel(validos, rechazados, metricas, ruta, fecha_operacion=None):
    """Genera un libro Excel con hojas ordenadas y un recuadro de estadisticas.

    Es un extra opcional. Si openpyxl no esta instalado, se lanza ImportError
    y el programa sigue funcionando con los CSV. `fecha_operacion` es la fecha
    que trae el JSON; se usa para que el reporte sea reproducible.
    """
    try:
        from openpyxl import Workbook
        from openpyxl.chart import BarChart, Reference
        from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
        from openpyxl.utils import get_column_letter
        from openpyxl.worksheet.properties import PageSetupProperties
    except ImportError:
        raise ImportError(
            "falta openpyxl. Instalalo con: python -m pip install openpyxl")

    azul = "1F4E78"
    azul_claro = "D9E2F3"
    azul_encabezado = "2F75B5"
    gris = "F2F2F2"
    dorado = "FFF2CC"
    verde = "E2EFDA"

    borde = Side(style="thin", color="BFBFBF")
    borde_medio = Side(style="medium", color=azul)

    fuente_titulo = Font(name="Calibri", size=16, bold=True, color="FFFFFF")
    fuente_seccion = Font(name="Calibri", size=11, bold=True, color=azul)
    fuente_encabezado = Font(name="Calibri", size=10, bold=True, color="FFFFFF")
    fuente_normal = Font(name="Calibri", size=10, color="000000")
    fuente_negrita = Font(name="Calibri", size=10, bold=True, color="000000")
    fuente_nota = Font(name="Calibri", size=9, italic=True, color="666666")

    relleno_titulo = PatternFill("solid", fgColor=azul)
    relleno_seccion = PatternFill("solid", fgColor=azul_claro)
    relleno_encabezado = PatternFill("solid", fgColor=azul_encabezado)
    relleno_gris = PatternFill("solid", fgColor=gris)
    relleno_dorado = PatternFill("solid", fgColor=dorado)
    relleno_verde = PatternFill("solid", fgColor=verde)
    relleno_banda = PatternFill("solid", fgColor="F7F9FC")

    centro = Alignment(horizontal="center", vertical="center", wrap_text=True)
    izquierda = Alignment(horizontal="left", vertical="center")
    derecha = Alignment(horizontal="right", vertical="center")

    def aplicar_borde(celda, izquierda_media=False, derecha_media=False,
                      arriba_media=False, abajo_media=False):
        celda.border = Border(
            left=borde_medio if izquierda_media else borde,
            right=borde_medio if derecha_media else borde,
            top=borde_medio if arriba_media else borde,
            bottom=borde_medio if abajo_media else borde,
        )

    def escribir_tabla(hoja, titulo, campos, filas, anchos, formatos,
                       colores=None):
        ultima_col = len(campos)
        ultima_letra = get_column_letter(ultima_col)

        hoja.merge_cells(start_row=1, start_column=1, end_row=1,
                         end_column=ultima_col)
        hoja.cell(row=1, column=1, value=titulo)
        for col in range(1, ultima_col + 1):
            celda = hoja.cell(row=1, column=col)
            celda.fill = relleno_titulo
            celda.font = fuente_titulo
            celda.alignment = centro
            aplicar_borde(celda, izquierda_media=col == 1,
                          derecha_media=col == ultima_col,
                          arriba_media=True, abajo_media=True)
        hoja.row_dimensions[1].height = 30

        for col, campo in enumerate(campos, start=1):
            celda = hoja.cell(row=2, column=col, value=campo)
            celda.fill = relleno_encabezado
            celda.font = fuente_encabezado
            celda.alignment = centro
            aplicar_borde(celda)
        hoja.row_dimensions[2].height = 24

        for indice, fila in enumerate(filas):
            numero_fila = 3 + indice
            for col, campo in enumerate(campos, start=1):
                valor = fila.get(campo, "")
                celda = hoja.cell(row=numero_fila, column=col, value=valor)
                celda.font = fuente_normal
                if campo in formatos:
                    celda.number_format = formatos[campo]
                if campo in ("entrada", "salida", "solicito_lavado",
                             "tope_aplicado"):
                    celda.alignment = centro
                elif isinstance(valor, (int, float)) and not isinstance(valor, bool):
                    celda.alignment = derecha
                else:
                    celda.alignment = izquierda
                if indice % 2 == 1:
                    celda.fill = relleno_banda
                if colores and indice < len(colores) and colores[indice]:
                    celda.fill = colores[indice]
                aplicar_borde(celda)

        for col, campo in enumerate(campos, start=1):
            hoja.column_dimensions[get_column_letter(col)].width = anchos.get(
                campo, 14)

        ultima_fila = max(2, 2 + len(filas))
        hoja.auto_filter.ref = "A2:{}{}".format(ultima_letra, ultima_fila)
        hoja.freeze_panes = "A3"
        hoja.sheet_view.showGridLines = False
        hoja.page_setup.orientation = "landscape"
        hoja.page_setup.fitToWidth = 1
        hoja.page_setup.fitToHeight = 0
        hoja.sheet_properties.pageSetUpPr = PageSetupProperties(fitToPage=True)
        hoja.print_title_rows = "1:2"

    fecha_reporte = datetime.now()
    fecha_texto = None
    if fecha_operacion:
        try:
            fecha_reporte = datetime.strptime(str(fecha_operacion), "%Y-%m-%d")
            fecha_texto = fecha_reporte.strftime("%d/%m/%Y")
        except ValueError:
            pass

    libro = Workbook()
    libro.properties.creator = "AutoPark"
    libro.properties.title = "Liquidación y estadísticas"
    libro.properties.created = fecha_reporte
    libro.properties.modified = fecha_reporte
    hoja_stats = libro.active
    hoja_stats.title = "Estadísticas"
    hoja_liq = libro.create_sheet("Liquidación")
    hoja_rank = libro.create_sheet("Ranking")
    hoja_rech = libro.create_sheet("Rechazados")

    # --- Hoja de liquidacion: ordenada por patente y luego por ID ---
    escribir_tabla(
        hoja_liq,
        "LIQUIDACIÓN DE MOVIMIENTOS VÁLIDOS",
        CAMPOS_LIQUIDACION,
        ordenar_liquidacion(validos),
        ANCHOS_LIQUIDACION,
        FORMATOS_LIQUIDACION,
    )
    hoja_liq.sheet_properties.tabColor = "2E75B6"

    # --- Hoja de ranking: de mayor a menor total, con puesto y top 3 resaltado ---
    filas_ranking = []
    for puesto, fila in enumerate(
            sorted(validos, key=lambda f: (-f["total"], f["patente"])), start=1):
        copia = dict(fila)
        copia["puesto"] = puesto
        filas_ranking.append(copia)

    colores_ranking = [relleno_dorado if i < 3 else None
                       for i in range(len(filas_ranking))]
    escribir_tabla(
        hoja_rank,
        "RANKING POR TOTAL COBRADO (MAYOR A MENOR)",
        CAMPOS_RANKING,
        filas_ranking,
        ANCHOS_LIQUIDACION,
        FORMATOS_LIQUIDACION,
        colores_ranking,
    )
    hoja_rank.sheet_properties.tabColor = "BF9000"

    # --- Hoja de rechazados: ordenada por ID ---
    escribir_tabla(
        hoja_rech,
        "MOVIMIENTOS RECHAZADOS",
        CAMPOS_RECHAZADOS,
        ordenar_rechazados(rechazados),
        ANCHOS_RECHAZADOS,
        {},
    )
    hoja_rech.sheet_properties.tabColor = "C00000"

    # --- Hoja de estadisticas: recuadro con el resumen del lote ---
    total_recaudado = sum(fila["total"] for fila in validos)
    total_movimientos = len(validos) + len(rechazados)
    ticket_promedio = total_recaudado / len(validos) if validos else 0
    mas_rentable = metricas["mas_rentable"] or "-"
    recaudado_mas_rentable = (
        metricas["recaudado"].get(mas_rentable, 0) if mas_rentable != "-" else 0)

    hoja_stats.merge_cells("A1:C1")
    hoja_stats["A1"] = "AUTOPARK — ESTADÍSTICAS DEL LOTE"
    hoja_stats["A1"].font = fuente_titulo
    hoja_stats["A1"].alignment = centro
    hoja_stats.row_dimensions[1].height = 32
    for col in range(1, 4):
        celda = hoja_stats.cell(row=1, column=col)
        celda.fill = relleno_titulo
        aplicar_borde(celda, izquierda_media=col == 1,
                      derecha_media=col == 3, arriba_media=True)

    hoja_stats.merge_cells("A2:C2")
    if fecha_texto:
        hoja_stats["A2"] = "Fecha de operación: {}".format(fecha_texto)
    elif fecha_operacion:
        hoja_stats["A2"] = "Fecha de operación: {}".format(fecha_operacion)
    else:
        hoja_stats["A2"] = "Reporte generado por AutoPark"
    hoja_stats["A2"].font = fuente_nota
    hoja_stats["A2"].alignment = derecha
    for col in range(1, 4):
        celda = hoja_stats.cell(row=2, column=col)
        celda.fill = PatternFill("solid", fgColor="FFFFFF")
        aplicar_borde(celda, izquierda_media=col == 1, derecha_media=col == 3)
    hoja_stats.row_dimensions[2].height = 16

    fila = 4

    def seccion(texto):
        nonlocal fila
        hoja_stats.merge_cells(start_row=fila, start_column=1,
                               end_row=fila, end_column=3)
        hoja_stats.cell(row=fila, column=1, value=texto)
        for col in range(1, 4):
            celda = hoja_stats.cell(row=fila, column=col)
            celda.fill = relleno_seccion
            celda.font = fuente_seccion
            celda.alignment = izquierda
            aplicar_borde(celda, izquierda_media=col == 1,
                          derecha_media=col == 3)
        hoja_stats.row_dimensions[fila].height = 20
        fila += 1

    def dato(etiqueta, valor, formato=None, negrita=False, relleno=None):
        nonlocal fila
        c1 = hoja_stats.cell(row=fila, column=1, value=etiqueta)
        c2 = hoja_stats.cell(row=fila, column=2, value=valor)
        c3 = hoja_stats.cell(row=fila, column=3, value="")
        for celda in (c1, c2, c3):
            celda.font = fuente_negrita if negrita else fuente_normal
            celda.alignment = izquierda
            if relleno:
                celda.fill = relleno
            aplicar_borde(celda, izquierda_media=celda.column == 1,
                          derecha_media=celda.column == 3)
        c2.alignment = derecha if isinstance(valor, (int, float)) else izquierda
        if formato:
            c2.number_format = formato
        fila += 1

    def tabla(fila_valores, formatos=None, negrita=False, relleno=None,
              encabezado=False):
        nonlocal fila
        for col, valor in enumerate(fila_valores, start=1):
            celda = hoja_stats.cell(row=fila, column=col, value=valor)
            celda.font = fuente_encabezado if encabezado else (
                fuente_negrita if negrita else fuente_normal)
            celda.alignment = centro if encabezado else (
                derecha if isinstance(valor, (int, float)) else izquierda)
            if encabezado:
                celda.fill = relleno_encabezado
            elif relleno:
                celda.fill = relleno
            if formatos and col - 1 < len(formatos) and formatos[col - 1]:
                celda.number_format = formatos[col - 1]
            aplicar_borde(celda, izquierda_media=col == 1,
                          derecha_media=col == 3)
        fila += 1

    seccion("RESUMEN GENERAL")
    dato("Total recaudado", total_recaudado, '"$"#,##0', negrita=True,
         relleno=relleno_verde)
    dato("Movimientos procesados", len(validos), "0")
    dato("Movimientos rechazados", len(rechazados), "0")
    dato("Total de movimientos", total_movimientos, "0")
    dato("Lavados realizados", metricas["lavados"], "0")
    dato("Porcentaje de lavados", metricas["porcentaje_lavado"] / 100, "0.0%")
    dato("Ticket promedio", ticket_promedio, '"$"#,##0')
    dato("Categoría más rentable", mas_rentable)
    dato("Recaudación de la más rentable", recaudado_mas_rentable,
         '"$"#,##0')

    fila += 1
    seccion("RECAUDACIÓN POR CATEGORÍA")
    tabla(["Categoría", "Recaudación", "Vehículos"], encabezado=True)
    fila_inicio_categorias = fila - 1
    categorias = sorted(
        metricas["recaudado"], key=lambda t: (-metricas["recaudado"][t], t))
    for categoria in categorias:
        tabla([categoria,
               metricas["recaudado"][categoria],
               metricas["cantidad"][categoria]],
              ['"$"#,##0', "0"])
    fila_fin_categorias = fila - 1
    tabla(["TOTAL", total_recaudado, len(validos)],
          ['"$"#,##0', "0"], negrita=True, relleno=relleno_gris)

    fila += 1
    seccion("ESTADÍA PROMEDIO POR CATEGORÍA")
    tabla(["Categoría", "Promedio (min)", "Equivalente"], encabezado=True)
    for categoria in sorted(metricas["promedios"]):
        promedio = metricas["promedios"][categoria]
        tabla([categoria, round(promedio, 1),
               formatear_minutos(int(round(promedio)))],
              ["0.0", None])

    fila += 1
    seccion("RECHAZOS POR MOTIVO")
    tabla(["Motivo", "Cantidad", ""], encabezado=True)
    rechazos_por_motivo = contar_rechazos(rechazados)
    if rechazos_por_motivo:
        for motivo, cantidad in rechazos_por_motivo.items():
            tabla([motivo, cantidad, ""], ["0", None])
    else:
        tabla(["Sin rechazos", 0, ""], ["0", None])

    fila += 1
    seccion("TOP 5 DEL RANKING")
    tabla(["Puesto", "Patente", "Total"], encabezado=True)
    top_5 = sorted(validos, key=lambda f: (-f["total"], f["patente"]))[:5]
    if top_5:
        for puesto, registro in enumerate(top_5, start=1):
            tabla([puesto, registro["patente"], registro["total"]],
                  ["0", '"$"#,##0'])
    else:
        tabla(["-", "-", 0], [None, '"$"#,##0'])

    # Recuadro exterior y formato de la hoja de estadisticas.
    ultima_fila_stats = hoja_stats.max_row
    for numero, fila_celdas in enumerate(
            hoja_stats.iter_rows(min_row=1, max_row=ultima_fila_stats,
                                 min_col=1, max_col=3), start=1):
        for celda in fila_celdas:
            aplicar_borde(
                celda,
                izquierda_media=celda.column == 1,
                derecha_media=celda.column == 3,
                arriba_media=numero == 1,
                abajo_media=numero == ultima_fila_stats,
            )

    hoja_stats.column_dimensions["A"].width = 34
    hoja_stats.column_dimensions["B"].width = 22
    hoja_stats.column_dimensions["C"].width = 20
    hoja_stats.sheet_view.showGridLines = False
    hoja_stats.page_setup.orientation = "landscape"
    hoja_stats.page_setup.fitToWidth = 1
    hoja_stats.page_setup.fitToHeight = 0
    hoja_stats.sheet_properties.pageSetUpPr = PageSetupProperties(fitToPage=True)
    hoja_stats.sheet_properties.tabColor = "375623"

    # Grafico de barras al costado del recuadro.
    if categorias:
        grafico = BarChart()
        grafico.type = "col"
        grafico.style = 10
        grafico.title = "Recaudación por categoría"
        grafico.y_axis.title = "Pesos"
        grafico.x_axis.title = "Categoría"
        grafico.y_axis.numFmt = '"$"#,##0'
        grafico.legend = None
        datos = Reference(hoja_stats, min_col=2,
                          min_row=fila_inicio_categorias,
                          max_row=fila_fin_categorias)
        etiquetas = Reference(hoja_stats, min_col=1,
                              min_row=fila_inicio_categorias + 1,
                              max_row=fila_fin_categorias)
        grafico.add_data(datos, titles_from_data=True)
        grafico.set_categories(etiquetas)
        grafico.height = 8.5
        grafico.width = 16
        hoja_stats.add_chart(grafico, "E4")

    libro.active = 0
    libro.save(ruta)
    return ruta
